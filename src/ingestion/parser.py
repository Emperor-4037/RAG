import os
from pathlib import Path
import pymupdf4llm
import docx
import pptx

def parse_document(file_path: Path) -> str:
    """
    Extract text from PDF, DOCX, or PPTX files.
    """
    ext = file_path.suffix.lower()
    
    if ext == ".pdf":
        # CPU-efficient markdown extraction (optionally handles OCR if configured)
        return pymupdf4llm.to_markdown(str(file_path))
        
    elif ext == ".docx":
        doc = docx.Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs])
        
    elif ext == ".pptx":
        prs = pptx.Presentation(file_path)
        text_blocks = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_blocks.append(shape.text)
        return "\n".join(text_blocks)
        
    else:
        # Fallback to raw text
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
